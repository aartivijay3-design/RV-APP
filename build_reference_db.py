# -*- coding: utf-8 -*-
"""
Builds data/reference_library.json — the BAWA house-style reference database —
from the Musterreiseverlaeufe sample itineraries (Word/PDF).

Parses each sample into hotel-description and sightseeing paragraph entries,
keyed by destination, so the RV app can reuse already-written Bawa house-style
language for a hotel or sight instead of asking the AI to invent it from
scratch every time (see reference_db.py for how app.py consumes this file).

Re-run this whenever new sample itineraries are added to the Muster folder:

    python build_reference_db.py

Override the source folder with the MUSTER_DIR env var if it ever moves.
"""
import json
import os
import re
import sys
import datetime
from pathlib import Path

import docx
import pypdf

DEFAULT_SRC_ROOT = (
    r"C:\Users\AartiVijay\OneDrive - BAWA Tours&Travel GmbH\BAWA - Bawa"
    r"\KUNDEN\Muster\Musterreiseverläufe"
)
SRC_ROOT = Path(os.environ.get("MUSTER_DIR", DEFAULT_SRC_ROOT))
OUT_PATH = Path(__file__).parent / "data" / "reference_library.json"

DAY_RE = re.compile(r"^tag\s*\d+", re.IGNORECASE)
HOTEL_RE = re.compile(r"^übernachtung\s+(?:im|in|bei)\s+(.+)$", re.IGNORECASE)
URL_RE = re.compile(r"^(https?://|www\.)", re.IGNORECASE)
END_RE = re.compile(r"^ende der reise$", re.IGNORECASE)
# Section headers introducing a bullet-style inclusions/exclusions list, not a
# place name. The generic "short line, no terminal punctuation" heading rule
# below would otherwise treat "Inkludierte Leistungen:" as a location heading
# and file the whole list that follows as if it were sightseeing prose for
# whatever place happened to be mentioned last — real corruption found in
# reference_library.json: an old Baltikum sample's full inclusions list (city
# tax, transfers, hotel nights, "Deutschsprachige Reiseleitung: in Litauen...")
# got stored as a "sightseeing" paragraph and later verbatim-injected into an
# unrelated day of a different Baltikum itinerary because a word in it
# happened to match that day's activities. Everything from one of these
# headers until the next day marker is skipped entirely instead.
LEISTUNGEN_HEADING_RE = re.compile(
    r"^(?:in|ex)?kludierte\s+leistungen\s*:?$"
    r"|^leistungen\s*:?$"
    r"|^optional\s*:?$"
    r"|^nicht\s+(?:inkludiert|enthalten|eingeschlossen)\s*:?$"
    r"|^ausgeschlossene\s+leistungen\s*:?$",
    re.IGNORECASE,
)
# Running page headers/footers in PDFs (e.g. "JAPAN – RUNDREISE") and bare page numbers
PAGE_NUMBER_RE = re.compile(r"^\d{1,4}$")
RUNNING_HEADER_RE = re.compile(r"^[A-ZÄÖÜß][A-ZÄÖÜß\s\-–—]{3,60}$")

STOPWORDS = {"hotel", "resort", "the", "le", "la", "des", "de", "und", "im",
             "in", "spa", "villa", "villas", "camp", "lodge", "house", "an", "am"}

# "alt" and "english" folders mix multiple countries per file — recovered from
# the filename instead of the (uninformative) folder name.
KNOWN_COUNTRIES = [
    "Bhutan", "Finnland", "Island", "Israel", "Japan", "Kambodscha", "Kanada",
    "Oman", "Schottland", "Taiwan", "Thailand", "Vietnam", "Greece", "China",
]


def normalize_ws(t: str) -> str:
    return re.sub(r"[ \t]+", " ", t).strip()


def clean_hotel_name(name: str) -> str:
    name = name.strip()
    # Strip a URL glued onto the same line (PDF extraction can merge them, and
    # font ligatures can mangle "http" itself — "://" survives intact though)
    name = re.split(r"\s*\S*://", name)[0]
    name = re.split(r"\s+www\.", name, flags=re.IGNORECASE)[0]
    name = re.sub(r"\s*(?:\bor\b|\boder\b)\s+.*$", "", name, flags=re.IGNORECASE)
    name = name.rstrip(" .,;:")
    return name


def significant_tokens(name: str):
    toks = re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]+", name)
    return [t for t in toks if len(t) > 3 and t.lower() not in STOPWORDS]


def paragraph_mentions_hotel(text: str, hotel_name: str, location_heading: str = "") -> bool:
    """True if `text` looks like it's describing the hotel itself, not just
    mentioning the city it's in. Tokens shared with the current location
    heading (e.g. hotel names like "Four Seasons Kyoto") are excluded, since
    those match almost any paragraph about that city, not just the hotel.
    """
    loc_low = location_heading.lower()
    toks = [t for t in significant_tokens(hotel_name) if t.lower() not in loc_low][:3]
    if not toks:
        return False
    low = text.lower()
    return any(re.search(rf"\b{re.escape(t.lower())}\b", low) for t in toks)


EN_MARKERS = re.compile(r"\b(the|you|your|and|with|please|arrival|departure|guide|will)\b", re.IGNORECASE)
DE_MARKERS = re.compile(r"\b(und|sie|ihre|ihren|der|die|das|für|mit|wird|besuchen|anschließend)\b", re.IGNORECASE)


def detect_language(lines) -> str:
    """Some 'alt' folder files are English-language client documents mixed
    into otherwise-German destination folders — detect from content, not
    just the folder name, so English text doesn't leak into the German
    house-style reference."""
    sample = " ".join(lines[:80])
    en_hits = len(EN_MARKERS.findall(sample))
    de_hits = len(DE_MARKERS.findall(sample))
    return "en" if en_hits > de_hits * 1.5 else "de"


def destination_for_file(folder_name: str, filename_stem: str) -> str:
    if folder_name not in ("alt", "english"):
        return folder_name
    for c in KNOWN_COUNTRIES:
        if filename_stem.lower().startswith(c.lower()):
            return c
    return filename_stem.split()[0]


def docx_lines_from_stream(stream):
    """Shared by docx_lines(path) and the app's document-upload-and-parse
    feature (which reads from an in-memory BytesIO, not a file on disk)."""
    d = docx.Document(stream)
    for p in d.paragraphs:
        t = normalize_ws(p.text)
        if not t:
            continue
        for sub in t.split("\n"):
            sub = sub.strip()
            if sub:
                yield sub
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = normalize_ws(p.text)
                    if t:
                        yield t


def docx_lines(path: Path):
    yield from docx_lines_from_stream(str(path))


def reflow_pdf_text(text: str):
    """pypdf yields one line per visual PDF line, chopping sentences at the
    line wrap. Rejoin wrapped lines into real paragraphs: merge a line into
    the running buffer unless the buffer already ends a sentence, or the
    line itself is a day/hotel marker (which always starts a new paragraph).
    """
    lines = [normalize_ws(l) for l in text.split("\n")]
    paragraphs = []
    buf = ""
    for line in lines:
        if not line:
            if buf:
                paragraphs.append(buf.strip())
                buf = ""
            continue
        if not buf:
            buf = line
            continue
        if DAY_RE.match(line) or HOTEL_RE.match(line):
            paragraphs.append(buf.strip())
            buf = line
            continue
        if buf.endswith("-"):
            buf = buf[:-1] + line
        elif buf.endswith((".", "!", "?", ":", ";")):
            paragraphs.append(buf.strip())
            buf = line
        else:
            buf = buf + " " + line
    if buf:
        paragraphs.append(buf.strip())
    return [p for p in paragraphs if p]


def pdf_lines_from_stream(stream):
    """Shared by pdf_lines(path) and the app's document-upload-and-parse
    feature (which reads from an in-memory BytesIO, not a file on disk)."""
    reader = pypdf.PdfReader(stream)
    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    for para in reflow_pdf_text(full_text):
        yield para


def pdf_lines(path: Path):
    yield from pdf_lines_from_stream(str(path))


def parse_lines(lines, destination: str, source_file: str):
    """Shared state machine for both docx and pdf line streams."""
    entries = []
    location_heading = ""
    bucket = []  # unclassified body paragraphs since the last location heading
    skip_section = False  # inside an "Inkludierte Leistungen"/"Optional" list

    def flush_as_sightseeing():
        for t in bucket:
            if len(t) >= 60:
                entries.append({
                    "type": "sightseeing",
                    "destination": destination,
                    "source_file": source_file,
                    "location_heading": location_heading,
                    "text": t,
                })
        bucket.clear()

    for text in lines:
        if URL_RE.match(text) or END_RE.match(text):
            continue
        if PAGE_NUMBER_RE.match(text) or RUNNING_HEADER_RE.match(text):
            continue

        day_m = DAY_RE.match(text)
        if day_m:
            skip_section = False
            flush_as_sightseeing()
            trailing = text[day_m.end():].strip(" -–—/\t")
            if trailing and len(trailing) <= 55:
                location_heading = trailing
            continue

        if LEISTUNGEN_HEADING_RE.match(text):
            flush_as_sightseeing()
            skip_section = True
            continue

        if skip_section:
            continue

        m = HOTEL_RE.match(text)
        if m:
            raw_name = clean_hotel_name(m.group(1))
            if raw_name:
                hotel_paras, remaining = [], []
                for t in bucket:
                    if len(t) >= 60 and paragraph_mentions_hotel(t, raw_name, location_heading):
                        hotel_paras.append(t)
                    else:
                        remaining.append(t)
                for t in remaining:
                    if len(t) >= 60:
                        entries.append({
                            "type": "sightseeing",
                            "destination": destination,
                            "source_file": source_file,
                            "location_heading": location_heading,
                            "text": t,
                        })
                if hotel_paras:
                    entries.append({
                        "type": "hotel",
                        "destination": destination,
                        "source_file": source_file,
                        "hotel_name": raw_name,
                        "text": " ".join(hotel_paras),
                    })
                bucket.clear()
            continue

        # Location heading candidate: short line, no terminal sentence punctuation
        if len(text) <= 55 and not text.endswith((".", "!", "?")):
            flush_as_sightseeing()
            location_heading = text
            continue

        bucket.append(text)

    flush_as_sightseeing()
    return entries


def main():
    if not SRC_ROOT.exists():
        print(f"Source folder not found: {SRC_ROOT}")
        print("Set MUSTER_DIR env var to point at the Musterreiseverläufe folder.")
        sys.exit(1)

    all_entries = []
    files = sorted(list(SRC_ROOT.rglob("*.docx")) + list(SRC_ROOT.rglob("*.pdf")))
    errors = []

    for path in files:
        folder = path.parent.name
        destination = destination_for_file(folder, path.stem)
        try:
            if path.suffix.lower() == ".docx":
                lines = list(docx_lines(path))
            else:
                lines = list(pdf_lines(path))
            language = "en" if folder == "english" else detect_language(lines)
            entries = parse_lines(lines, destination, path.name)
            for e in entries:
                e["language"] = language
            all_entries.extend(entries)
        except Exception as e:
            errors.append(f"{path}: {e!r}")

    hotel_count = sum(1 for e in all_entries if e["type"] == "hotel")
    sight_count = sum(1 for e in all_entries if e["type"] == "sightseeing")
    destinations = sorted(set(e["destination"] for e in all_entries))

    out = {
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "source_root": str(SRC_ROOT),
        "source_file_count": len(files),
        "stats": {
            "total_entries": len(all_entries),
            "hotel_entries": hotel_count,
            "sightseeing_entries": sight_count,
            "destinations": destinations,
        },
        "entries": all_entries,
    }

    out_path = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT_PATH
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Files processed: {len(files)}  (errors: {len(errors)})")
    print(f"Entries: {len(all_entries)}  (hotel: {hotel_count}, sightseeing: {sight_count})")
    print(f"Destinations ({len(destinations)}): {', '.join(destinations)}")
    if errors:
        print("\nErrors:")
        for e in errors:
            print(" -", e)
    print(f"\nWritten to {out_path}")


if __name__ == "__main__":
    main()
